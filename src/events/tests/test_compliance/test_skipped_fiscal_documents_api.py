"""Owner-facing views of skipped fiscal documents: list, resolve, ticket flag, draft hint, report (#1091)."""

import datetime as dt
import io
import typing as t
from decimal import Decimal
from unittest.mock import patch

import orjson
import pytest
from django.db import connection
from django.test.client import Client
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from ninja_jwt.tokens import RefreshToken
from openpyxl import load_workbook

from accounts.models import RevelUser
from events.models import (
    Event,
    Organization,
    OrganizationStaff,
    Payment,
    PermissionMap,
    PermissionsSchema,
    SkippedFiscalDocument,
    TicketTier,
)
from events.service import revenue_report_service as report
from events.service.attendee_invoice_service import generate_attendee_invoice
from events.tests.test_attendee_invoice._helpers import (
    MOCK_RENDER_PDF,
    _create_draft,
    _create_payment,
    _default_billing_snapshot,
    _make_org_invoicing_ready,
)

pytestmark = [pytest.mark.django_db, pytest.mark.usefixtures("_no_pdf")]


@pytest.fixture
def _no_pdf() -> t.Iterator[None]:
    with patch(MOCK_RENDER_PDF, return_value=b"fake-pdf"):
        yield


def _ready_in(org: Organization, country: str) -> Organization:
    _make_org_invoicing_ready(org)
    org.vat_country_code = country
    org.vat_id = f"{country}123456789"
    org.invoicing_mode = Organization.InvoicingMode.AUTO
    org.save()
    return org


def _be_business() -> dict[str, t.Any]:
    return {**_default_billing_snapshot(), "vat_id": "BE0123456789", "vat_country_code": "BE"}


def _skip_sale(event: Event, tier: TicketTier, user: RevelUser, session_id: str) -> Payment:
    """A BE domestic-B2B sale whose invoice generation is skipped and recorded."""
    payment = _create_payment(
        user=user, event=event, tier=tier, session_id=session_id, buyer_billing_snapshot=_be_business()
    )
    assert generate_attendee_invoice(session_id) is None
    return payment


@pytest.fixture
def skipped(
    organization: Organization, event: Event, event_ticket_tier: TicketTier, member_user: RevelUser
) -> SkippedFiscalDocument:
    _ready_in(organization, "BE")
    _skip_sale(event, event_ticket_tier, member_user, "cs_skip_1")
    return SkippedFiscalDocument.objects.get()


def _client_for(user: RevelUser) -> Client:
    refresh = RefreshToken.for_user(user)
    return Client(HTTP_AUTHORIZATION=f"Bearer {str(refresh.access_token)}")  # type: ignore[attr-defined]


@pytest.fixture
def ticket_staff_client(organization: Organization, organization_staff_user: RevelUser) -> Client:
    """Staff with every ticket permission: still not allowed on the owner-only endpoints."""
    OrganizationStaff.objects.create(
        organization=organization,
        user=organization_staff_user,
        permissions=PermissionsSchema(default=PermissionMap(manage_tickets=True, manage_event=True)).model_dump(
            mode="json"
        ),
    )
    return _client_for(organization_staff_user)


def _list_url(org: Organization) -> str:
    return reverse("api:list_skipped_fiscal_documents", kwargs={"slug": org.slug})


def _resolve_url(org: Organization, doc: SkippedFiscalDocument) -> str:
    return reverse("api:resolve_skipped_fiscal_document", kwargs={"slug": org.slug, "document_id": doc.id})


class TestList:
    def test_owner_sees_the_document(
        self, owner_client: Client, organization: Organization, event: Event, skipped: SkippedFiscalDocument
    ) -> None:
        response = owner_client.get(_list_url(organization))

        assert response.status_code == 200
        body = response.json()
        assert body["count"] == 1
        row = body["results"][0]
        assert row["id"] == str(skipped.id)
        assert row["kind"] == "invoice"
        assert row["reason_code"] == "b2b_e_invoicing"
        assert row["policy_country"] == "BE"
        assert "Peppol" in row["reason"]
        assert row["event_id"] == str(event.id) and row["event_name"] == event.name
        assert row["buyer_vat_id"] == "BE0123456789"
        assert row["total_gross"] == "100.00"
        assert row["vat_breakdown"][0]["vat_rate"] == "22.00"
        assert row["ticket_ids"] == [str(skipped.payments.get().ticket_id)]
        assert row["invoice_number"] is None and row["resolved_at"] is None

    def test_filters(
        self,
        owner_client: Client,
        organization: Organization,
        event: Event,
        event_ticket_tier: TicketTier,
        member_user: RevelUser,
        skipped: SkippedFiscalDocument,
    ) -> None:
        _skip_sale(event, event_ticket_tier, member_user, "cs_skip_2")
        SkippedFiscalDocument.objects.filter(stripe_session_id="cs_skip_2").update(resolved_at=skipped.decided_at)

        def ids(**params: t.Any) -> list[str]:
            return [r["id"] for r in owner_client.get(_list_url(organization), params).json()["results"]]

        assert ids(resolved="false") == [str(skipped.id)]
        assert len(ids(resolved="true")) == 1
        assert len(ids(kind="invoice")) == 2
        assert ids(kind="credit_note") == []
        assert ids(reason_code="fiscalized_invoicing") == []
        assert len(ids(event_id=str(event.id))) == 2
        assert len(ids(search="buyer@example.de")) == 2

    def test_other_organizations_documents_are_not_listed(
        self, owner_client: Client, organization: Organization, skipped: SkippedFiscalDocument
    ) -> None:
        SkippedFiscalDocument.objects.update(organization=None)

        assert owner_client.get(_list_url(organization)).json()["count"] == 0

    def test_staff_is_forbidden(
        self, ticket_staff_client: Client, organization: Organization, skipped: SkippedFiscalDocument
    ) -> None:
        assert ticket_staff_client.get(_list_url(organization)).status_code == 403
        response = ticket_staff_client.post(
            _resolve_url(organization, skipped),
            data=orjson.dumps({"external_reference": "X"}),
            content_type="application/json",
        )
        assert response.status_code == 403


class TestResolve:
    def test_owner_resolves_with_a_reference(
        self,
        owner_client: Client,
        organization: Organization,
        organization_owner_user: RevelUser,
        skipped: SkippedFiscalDocument,
    ) -> None:
        response = owner_client.post(
            _resolve_url(organization, skipped),
            data=orjson.dumps({"external_reference": "  PEPPOL-2026-0042 "}),
            content_type="application/json",
        )

        assert response.status_code == 200
        assert response.json()["external_reference"] == "PEPPOL-2026-0042"
        assert response.json()["resolved_at"] is not None
        skipped.refresh_from_db()
        assert skipped.resolved_by == organization_owner_user
        assert skipped.external_reference == "PEPPOL-2026-0042"

    def test_document_of_a_deleted_event_and_buyer_can_still_be_resolved(
        self,
        owner_client: Client,
        organization: Organization,
        event: Event,
        member_user: RevelUser,
        skipped: SkippedFiscalDocument,
    ) -> None:
        """The SET_NULL foreign keys must not trip ``full_clean`` on save."""
        event.delete()
        member_user.delete()
        skipped.refresh_from_db()
        assert skipped.event is None and skipped.user is None
        assert skipped.buyer_vat_id == "BE0123456789"  # the snapshot outlives the cascaded payment

        response = owner_client.post(
            _resolve_url(organization, skipped),
            data=orjson.dumps({"external_reference": "EXT-9"}),
            content_type="application/json",
        )

        assert response.status_code == 200
        assert response.json()["event_name"] is None

    def test_blank_reference_is_rejected(
        self, owner_client: Client, organization: Organization, skipped: SkippedFiscalDocument
    ) -> None:
        response = owner_client.post(
            _resolve_url(organization, skipped),
            data=orjson.dumps({"external_reference": "   "}),
            content_type="application/json",
        )

        assert response.status_code == 422
        skipped.refresh_from_db()
        assert skipped.resolved_at is None

    def test_another_organizations_document_is_404(
        self, owner_client: Client, organization: Organization, skipped: SkippedFiscalDocument
    ) -> None:
        SkippedFiscalDocument.objects.update(organization=None)

        response = owner_client.post(
            _resolve_url(organization, skipped),
            data=orjson.dumps({"external_reference": "X"}),
            content_type="application/json",
        )

        assert response.status_code == 404


class TestTicketFlag:
    def _tickets(self, client: Client, event: Event, **params: t.Any) -> list[dict[str, t.Any]]:
        url = reverse("api:list_tickets", kwargs={"event_id": event.pk})
        response = client.get(url, {"include_past": "true", **params})
        assert response.status_code == 200
        results: list[dict[str, t.Any]] = response.json()["results"]
        return results

    def test_buyer_dashboard_does_not_offer_the_organizer_filter(
        self, event: Event, member_user: RevelUser, skipped: SkippedFiscalDocument
    ) -> None:
        """``invoice_skipped`` is admin-only: the buyer's ticket list ignores it rather than 500."""
        response = _client_for(member_user).get(
            reverse("api:dashboard_tickets"), {"invoice_skipped": "true", "include_past": "true"}
        )

        assert response.status_code == 200
        assert "invoice_skipped" not in response.json()["results"][0]

    def test_list_flags_and_filters_skipped_sales(
        self,
        ticket_staff_client: Client,
        event: Event,
        event_ticket_tier: TicketTier,
        member_user: RevelUser,
        skipped: SkippedFiscalDocument,
    ) -> None:
        """``manage_tickets`` staff see the flag (and can filter on it) without the owner-only list."""
        invoiced = _create_payment(
            user=member_user, event=event, tier=event_ticket_tier, session_id="cs_ok", buyer_billing_snapshot=None
        )
        skipped_ticket_id = str(skipped.payments.get().ticket_id)

        flags = {r["id"]: r["invoice_skipped"] for r in self._tickets(ticket_staff_client, event)}
        assert flags == {skipped_ticket_id: True, str(invoiced.ticket_id): False}
        assert [r["id"] for r in self._tickets(ticket_staff_client, event, invoice_skipped="true")] == [
            skipped_ticket_id
        ]
        assert [r["id"] for r in self._tickets(ticket_staff_client, event, invoice_skipped="false")] == [
            str(invoiced.ticket_id)
        ]

    def test_flag_adds_no_per_row_queries(
        self,
        owner_client: Client,
        event: Event,
        event_ticket_tier: TicketTier,
        member_user: RevelUser,
        skipped: SkippedFiscalDocument,
    ) -> None:
        with CaptureQueriesContext(connection) as one:
            self._tickets(owner_client, event)
        for i in range(3):
            _skip_sale(event, event_ticket_tier, member_user, f"cs_more_{i}")
        with CaptureQueriesContext(connection) as four:
            rows = self._tickets(owner_client, event)

        assert len(rows) == 4 and all(r["invoice_skipped"] for r in rows)
        assert len(four) == len(one)

    def test_single_ticket_detail_carries_the_flag(
        self, owner_client: Client, event: Event, skipped: SkippedFiscalDocument
    ) -> None:
        ticket_id = skipped.payments.get().ticket_id
        url = reverse("api:get_ticket", kwargs={"event_id": event.pk, "ticket_id": ticket_id})

        assert owner_client.get(url).json()["invoice_skipped"] is True


class TestDraftIssueBlockedReason:
    def _invoices(self, client: Client, org: Organization) -> list[dict[str, t.Any]]:
        url = reverse("api:list_attendee_invoices", kwargs={"slug": org.slug})
        results: list[dict[str, t.Any]] = client.get(url).json()["results"]
        return results

    def test_blocked_draft_explains_why_and_list_costs_no_per_row_queries(
        self, owner_client: Client, organization: Organization, event: Event, member_user: RevelUser
    ) -> None:
        _ready_in(organization, "HR")
        _create_draft(organization, event, member_user)
        with CaptureQueriesContext(connection) as one:
            rows = self._invoices(owner_client, organization)
        assert "Croatia" in rows[0]["issue_blocked_reason"]

        for _ in range(3):
            _create_draft(organization, event, member_user)
        with CaptureQueriesContext(connection) as four:
            rows = self._invoices(owner_client, organization)

        assert len(rows) == 4 and all(r["issue_blocked_reason"] for r in rows)
        assert len(four) == len(one)

    def test_issuable_draft_and_detail(
        self, owner_client: Client, organization: Organization, event: Event, member_user: RevelUser
    ) -> None:
        _ready_in(organization, "IT")
        draft = _create_draft(organization, event, member_user)
        url = reverse("api:get_attendee_invoice", kwargs={"slug": organization.slug, "invoice_id": draft.id})

        assert self._invoices(owner_client, organization)[0]["issue_blocked_reason"] == ""
        assert owner_client.get(url).json()["issue_blocked_reason"] == ""

    def test_b2b_draft_uses_the_checkout_vies_outcome(
        self,
        owner_client: Client,
        organization: Organization,
        event: Event,
        event_ticket_tier: TicketTier,
        member_user: RevelUser,
    ) -> None:
        """A Belgian business buyer's pre-gate draft is blocked; a VIES-rejected ID is a consumer's."""
        _ready_in(organization, "BE")
        for session_id, status in (("cs_valid", "valid"), ("cs_invalid", "invalid")):
            _create_payment(
                user=member_user,
                event=event,
                tier=event_ticket_tier,
                session_id=session_id,
                buyer_billing_snapshot={**_be_business(), "vat_id_status": status},
            )
            draft = _create_draft(organization, event, member_user)
            draft.stripe_session_id = session_id
            draft.buyer_vat_id = "BE0123456789"
            draft.save(update_fields=["stripe_session_id", "buyer_vat_id"])

        reasons = {r["id"]: r["issue_blocked_reason"] for r in self._invoices(owner_client, organization)}
        by_session = {d.stripe_session_id: reasons[str(d.id)] for d in organization.attendee_invoices.all()}
        assert "Peppol" in by_session["cs_valid"]
        assert by_session["cs_invalid"] == ""


def _wide(org: Organization, event_id: t.Any = None) -> report.ReportScope:
    return report.ReportScope(org=org, event_id=event_id, date_from=dt.date(2000, 1, 1), date_to=dt.date(2100, 1, 1))


class TestRevenueReport:
    def test_sheet_lists_skipped_documents_in_the_period(
        self, organization: Organization, skipped: SkippedFiscalDocument
    ) -> None:
        wb = load_workbook(io.BytesIO(report.build_xlsx(report.build_revenue_report_data(_wide(organization)))))

        sheet = wb["Invoices to issue yourself"]
        headers = [c.value for c in sheet[1]]
        row = dict(zip(headers, [c.value for c in sheet[2]], strict=False))
        assert row["document"] == "invoice"
        assert row["reason"] == "b2b_e_invoicing"
        assert row["policy_country"] == "BE"
        assert row["buyer_vat_id"] == "BE0123456789"
        assert row["vat_rates"] == "22%"
        assert row["gross"] == 100
        assert row["stripe_session_id"] == "cs_skip_1"
        # The explanatory note sits below the data, after one blank row.
        assert str(sheet.cell(row=4, column=1).value).startswith("Revel did not issue")

        outside = report.ReportScope(
            org=organization, event_id=None, date_from=dt.date(2000, 1, 1), date_to=dt.date(2000, 12, 31)
        )
        assert report.build_revenue_report_data(outside).skipped_documents == []

    def test_transactions_carry_buyer_country_and_payment_intent(
        self, organization: Organization, skipped: SkippedFiscalDocument
    ) -> None:
        """``buyer_country`` reads the snapshot's ``vat_country_code`` (was the missing ``country`` key)."""
        Payment.objects.update(stripe_payment_intent_id="pi_123")

        section = report.build_revenue_report_data(_wide(organization)).sections[0]

        txn = section.transactions[0]
        assert (txn.buyer_country, txn.stripe_payment_intent_id) == ("BE", "pi_123")

    def test_hash_changes_when_a_document_is_skipped_or_resolved(
        self,
        organization: Organization,
        event: Event,
        event_ticket_tier: TicketTier,
        member_user: RevelUser,
        organization_owner_user: RevelUser,
    ) -> None:
        _ready_in(organization, "BE")
        _create_payment(
            user=member_user,
            event=event,
            tier=event_ticket_tier,
            session_id="cs_skip_1",
            buyer_billing_snapshot=_be_business(),
        )
        before = report.compute_revenue_data_hash(_wide(organization))

        generate_attendee_invoice("cs_skip_1")
        skipped_hash = report.compute_revenue_data_hash(_wide(organization))
        assert skipped_hash != before

        from events.service.skipped_fiscal_document_service import resolve_skipped_document

        resolve_skipped_document(SkippedFiscalDocument.objects.get(), organization_owner_user, "EXT-1")
        assert report.compute_revenue_data_hash(_wide(organization)) != skipped_hash

    def test_event_scoped_report_only_lists_that_events_documents(
        self, organization: Organization, event: Event, skipped: SkippedFiscalDocument
    ) -> None:
        other = Event.objects.create(
            organization=organization, name="Other", slug="other", start=event.start, end=event.end
        )

        assert report.build_revenue_report_data(_wide(organization, event.id)).skipped_documents == [skipped]
        assert report.build_revenue_report_data(_wide(organization, other.id)).skipped_documents == []
        assert skipped.total_gross == Decimal("100.00")
